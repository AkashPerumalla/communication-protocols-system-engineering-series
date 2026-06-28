from pathlib import Path

from docx import Document as DocxDocument
from fpdf import FPDF


class SampleDocumentsService:
    def ensure_seed_documents(self, documents_dir: Path) -> None:
        documents_dir.mkdir(parents=True, exist_ok=True)

        txt_docs = {
            "device_manual_router.txt": "Router-01 Device Manual\nInterface monitoring interval is 30 seconds.\nAlarm threshold for packet loss is 2 percent.",
            "alarm_procedures.txt": "Alarm Procedures\nCritical alarm acknowledgment SLA is 5 minutes.\nEscalate to L2 after two failed remediation attempts.",
            "incident_response_runbook.txt": "Incident Response Runbook\nStep 1 isolate impacted node.\nStep 2 verify telemetry drift.\nStep 3 trigger rollback.",
            "telecom_sop.txt": "Telecom SOP\nUse maintenance window 02:00-04:00 UTC.\nLog all change request IDs in ticket notes.",
        }
        md_docs = {
            "network_troubleshooting_guide.md": "# Network Troubleshooting Guide\nCheck optical power and interface CRC counters before reboot actions.",
            "snmp_documentation.md": "# SNMP Documentation\nUse GET for direct OID reads, GETNEXT for tree walking, GETBULK for table retrieval.",
            "rest_api_documentation.md": "# REST API Documentation\nPlatform endpoint /api/platform/status returns health and model metadata.",
        }
        docx_docs = {
            "kubernetes_notes.docx": "Kubernetes Notes\nReadiness probes must validate application dependencies before traffic is accepted.",
            "docker_guide.docx": "Docker Guide\nUse multi-stage builds to reduce runtime image size and attack surface.",
        }
        pdf_docs = {
            "satcom_operations_guide.pdf": "SatCom Operations Guide\nTrack carrier lock, Eb/No, and modem BER metrics every minute.",
        }

        for name, content in txt_docs.items():
            (documents_dir / name).write_text(content, encoding="utf-8")

        for name, content in md_docs.items():
            (documents_dir / name).write_text(content, encoding="utf-8")

        for name, content in docx_docs.items():
            self._write_docx(documents_dir / name, content)

        for name, content in pdf_docs.items():
            self._write_pdf(documents_dir / name, content)

    def _write_docx(self, path: Path, content: str) -> None:
        doc = DocxDocument()
        for line in content.split("\n"):
            doc.add_paragraph(line)
        doc.save(str(path))

    def _write_pdf(self, path: Path, content: str) -> None:
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        for line in content.split("\n"):
            pdf.cell(w=180, h=8, text=line, new_x="LMARGIN", new_y="NEXT")
        pdf.output(str(path))
